import os
import re
import json
import unicodedata
from io import BytesIO
import cloudinary  # type: ignore
import cloudinary.uploader  # type: ignore
from datetime import datetime, timedelta

from flask import Flask, render_template, request, redirect, url_for, jsonify, flash, send_file  # type: ignore
from flask_login import (  # type: ignore
    LoginManager, UserMixin, login_user, logout_user,
    login_required
)
from sqlalchemy import create_engine, Column, Integer, String, Float, Date, DateTime, text, func, inspect, or_  # type: ignore
from sqlalchemy.orm import sessionmaker, declarative_base  # type: ignore
import pandas as pd  # type: ignore
from werkzeug.utils import secure_filename  # type: ignore
import math
from openpyxl import load_workbook  # type: ignore


# ---------- Config ----------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "shop.db")

# Local development secrets are kept outside Git and loaded before DB setup.
LOCAL_ENV_PATH = os.path.join(BASE_DIR, ".env.local")
if os.path.exists(LOCAL_ENV_PATH):
    with open(LOCAL_ENV_PATH, encoding="utf-8") as env_file:
        for env_line in env_file:
            env_line = env_line.strip()
            if not env_line or env_line.startswith("#") or "=" not in env_line:
                continue
            env_key, env_value = env_line.split("=", 1)
            env_key = env_key.strip()
            if not os.environ.get(env_key):
                os.environ[env_key] = env_value.strip()

ADMIN_USERNAME = os.environ.get("ADMIN_USERNAME", "admin")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "bike123")
SECRET_KEY = os.environ.get("SECRET_KEY", "doi-chuoi-bi-mat-nay-truoc-khi-deploy")
CLOUDINARY_CLOUD_NAME = os.environ.get("CLOUDINARY_CLOUD_NAME", "")
CLOUDINARY_API_KEY = os.environ.get("CLOUDINARY_API_KEY", "")
CLOUDINARY_API_SECRET = os.environ.get("CLOUDINARY_API_SECRET", "")

cloudinary.config(
    cloud_name=CLOUDINARY_CLOUD_NAME,
    api_key=CLOUDINARY_API_KEY,
    api_secret=CLOUDINARY_API_SECRET,
    secure=True,
)
app = Flask(__name__)
app.config["SECRET_KEY"] = SECRET_KEY


# ---------- Database ----------
Base = declarative_base()
DATABASE_URL = os.environ.get("DATABASE_URL", "").strip()
if DATABASE_URL:
    if DATABASE_URL.startswith("postgres://"):
        DATABASE_URL = DATABASE_URL.replace("postgres://", "postgresql://", 1)
    engine = create_engine(
        DATABASE_URL,
        pool_pre_ping=True,
        pool_recycle=300,
    )
else:
    engine = create_engine(
        f"sqlite:///{DB_PATH}",
        connect_args={"check_same_thread": False}
    )
SessionLocal = sessionmaker(bind=engine)
print(f"[APP DEBUG] Đang dùng DB: {engine.url}")

class Product(Base):
    __tablename__ = "products"

    id = Column(Integer, primary_key=True)
    code = Column(String, unique=True, nullable=False)
    model = Column(String, nullable=False)
    color = Column(String, default="")
    category = Column(String, default="")
    unit = Column(String, default="Chiếc")
    stock = Column(Integer, default=0)
    reserved = Column(Integer, default=0)
    price = Column(Float, nullable=True)
    retail_price = Column(Float, nullable=True)
    image_filename = Column(String, default="")
    updated_at = Column(DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)

    @property
    def available(self):
        stock = self.stock or 0
        reserved = self.reserved or 0
        return max(stock - reserved, 0)

    def to_dict(self):
        return {
            "id": self.id,
            "code": self.code,
            "model": self.model,
            "color": self.color,
            "category": self.category,
            "unit": self.unit,
            "stock": self.stock or 0,
            "reserved": self.reserved or 0,
            "available": self.available,
            "price": self.price,
            "image": self.image_filename,
        }


class AppSetting(Base):
    __tablename__ = "app_settings"

    key = Column(String, primary_key=True)
    value = Column(String, nullable=False)


class Supplier(Base):
    __tablename__ = "suppliers"

    id = Column(Integer, primary_key=True)
    code = Column(String, unique=True, nullable=False)
    name = Column(String, nullable=False)
    phone = Column(String, default="")
    email = Column(String, default="")
    address = Column(String, default="")
    created_at = Column(DateTime, default=datetime.utcnow)


class PurchaseReceipt(Base):
    __tablename__ = "purchase_receipts"

    id = Column(Integer, primary_key=True)
    receipt_number = Column(String, unique=True, nullable=False)
    supplier_id = Column(Integer, nullable=True)
    supplier_name = Column(String, nullable=False)
    received_at = Column(DateTime, default=datetime.utcnow)
    expected_arrival_date = Column(Date, nullable=True)
    notes = Column(String, default="")
    total_amount = Column(Float, default=0)
    created_at = Column(DateTime, default=datetime.utcnow)


class PurchaseReceiptItem(Base):
    __tablename__ = "purchase_receipt_items"

    id = Column(Integer, primary_key=True)
    receipt_id = Column(Integer, nullable=False)
    product_id = Column(Integer, nullable=False)
    product_code = Column(String, nullable=False)
    product_name = Column(String, nullable=False)
    color = Column(String, default="")
    unit = Column(String, default="Chiếc")
    quantity = Column(Integer, nullable=False)
    unit_price = Column(Float, default=0)
    line_total = Column(Float, default=0)


Base.metadata.create_all(engine)


def ensure_retail_price_column():
    """Add the retail price to existing SQLite and PostgreSQL databases."""
    columns = {column["name"] for column in inspect(engine).get_columns("products")}
    if "retail_price" not in columns:
        with engine.begin() as conn:
            conn.execute(text("ALTER TABLE products ADD COLUMN retail_price FLOAT"))


ensure_retail_price_column()


# ---------- Upload folder ----------
IMAGE_UPLOAD_DIR = os.path.join(BASE_DIR, "static", "uploads")
os.makedirs(IMAGE_UPLOAD_DIR, exist_ok=True)
SALES_HISTORY_FILE = os.path.join(BASE_DIR, "data", "san_pham_da_ban_12_thang.xlsm")


def ensure_columns():
    """
    Đảm bảo database cũ vẫn chạy được nếu trước đây thiếu cột.
    """
    with engine.connect() as conn:
        res = conn.execute(text("PRAGMA table_info(products)")).fetchall()
        cols = [r[1] for r in res]

        if "image_filename" not in cols:
            conn.execute(text("ALTER TABLE products ADD COLUMN image_filename VARCHAR"))

        if "unit" not in cols:
            conn.execute(text("ALTER TABLE products ADD COLUMN unit VARCHAR DEFAULT 'Chiếc'"))

        if "price" not in cols:
            conn.execute(text("ALTER TABLE products ADD COLUMN price FLOAT"))

        conn.commit()


if engine.dialect.name == "sqlite":
    try:
        ensure_columns()
    except Exception:
        pass


# ---------- Auth ----------
login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "admin_login"


class AdminUser(UserMixin):
    id = "admin"


@login_manager.user_loader
def load_user(user_id):
    if user_id == "admin":
        return AdminUser()
    return None


# ---------- Helpers ----------
def normalize_model_name(model_name: str):
    """
    Gộp các mẫu xe có phần ghi chú trong ngoặc về cùng một tên mẫu.

    Ví dụ:
    24BIKE Nhôm (tem đậm) -> 24BIKE Nhôm
    24BIKE Nhôm (tem lợt) -> 24BIKE Nhôm
    24BIKE Nhôm (tem xám) -> 24BIKE Nhôm
    """
    if not model_name:
        return ""

    name = str(model_name).strip()

    # Xóa phần ghi chú cuối cùng nằm trong ngoặc
    name = re.sub(r"\s*\([^)]*\)\s*$", "", name)

    # Xóa khoảng trắng dư
    name = re.sub(r"\s+", " ", name).strip()

    return name


def product_match_key(value):
    """Build a conservative key shared by Excel sales names and product models."""
    normalized = normalize_model_name(str(value or "")).casefold()
    normalized = unicodedata.normalize("NFD", normalized)
    normalized = "".join(
        char for char in normalized
        if unicodedata.category(char) != "Mn"
    )
    return re.sub(r"[^a-z0-9]+", "", normalized)


def load_monthly_sales():
    """Read net sales from the bundled 12-month workbook, grouped by model."""
    if not os.path.exists(SALES_HISTORY_FILE):
        return {}

    workbook = load_workbook(
        SALES_HISTORY_FILE,
        read_only=True,
        data_only=True,
        keep_vba=True,
    )
    monthly_sales = {}
    try:
        sheet = workbook[workbook.sheetnames[0]]
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_number < 7 or not row or not row[0]:
                continue

            key = product_match_key(row[0])
            if not key:
                continue

            sold_12_months = max(safe_float(row[7] if len(row) > 7 else 0) or 0, 0)
            monthly_sales[key] = monthly_sales.get(key, 0) + sold_12_months / 12
    finally:
        workbook.close()

    return monthly_sales


def split_model_color(name: str):
    """
    Tách dữ liệu nhập từ file dạng:
    'Tên xe - Màu' -> model, color

    Sau khi tách màu, model cũng được chuẩn hóa để tránh tách card ngoài trang khách.
    """
    if not name:
        return "", ""

    name = str(name).strip()
    parts = name.rsplit("-", 1)

    if len(parts) == 2 and parts[1].strip():
        model = normalize_model_name(parts[0].strip())
        color = parts[1].strip()
        return model, color

    return normalize_model_name(name), ""


def safe_int(value, default=0):
    try:
        if value is None or value == "":
            return default
        return int(float(value))
    except Exception:
        return default



def safe_float(value):
    try:
        if value is None or value == "":
            return None
        result = float(value)
        if math.isnan(result) or math.isinf(result):
            return None
        return result
    except Exception:
        return None


def save_product_image(file, product_id):
    if not file or not file.filename:
        return ""

    try:
        result = cloudinary.uploader.upload(
            file,
            folder="abraham_inventory/products",
            public_id=f"product_{product_id}",
            overwrite=True,
            resource_type="image",
        )
        return result.get("secure_url", "")
    except Exception as e:
        print(f"Loi upload Cloudinary: {e}")
        return ""


def seed_if_empty():
    db = SessionLocal()
    try:
        if db.query(Product).count() == 0:
            seed_path = os.path.join(BASE_DIR, "seed_data.json")
            if os.path.exists(seed_path):
                import json

                with open(seed_path, encoding="utf-8") as f:
                    rows = json.load(f)

                for r in rows:
                    db.add(Product(
                        code=r["code"],
                        model=normalize_model_name(r["model"]),
                        color=r.get("color", ""),
                        category=r.get("category", ""),
                        unit=r.get("unit", "Chiếc") or "Chiếc",
                        stock=safe_int(r.get("stock")),
                        reserved=safe_int(r.get("reserved")),
                        price=safe_float(r.get("price")),
                    ))

                db.commit()
    finally:
        db.close()


# ---------- Public routes ----------
def retail_prices_visible(db):
    setting = db.query(AppSetting).filter_by(key="retail_prices_visible").first()
    return setting is not None and setting.value == "1"


def render_catalog_page(price_type):
    db = SessionLocal()
    try:
        visible = retail_prices_visible(db)
        incoming_rows = (
            db.query(PurchaseReceiptItem, PurchaseReceipt.expected_arrival_date)
            .join(PurchaseReceipt, PurchaseReceipt.id == PurchaseReceiptItem.receipt_id)
            .filter(PurchaseReceipt.expected_arrival_date.isnot(None))
            .filter(PurchaseReceipt.expected_arrival_date > datetime.now().date())
            .order_by(PurchaseReceipt.expected_arrival_date.asc(), PurchaseReceiptItem.id.asc())
            .limit(30)
            .all()
        )
        incoming_products = [
            {
                "name": item.product_name,
                "date": expected_date.strftime("%d/%m/%Y"),
            }
            for item, expected_date in incoming_rows
        ]
    finally:
        db.close()
    return render_template(
        "catalog.html",
        price_type=price_type,
        price_label="Giá lẻ" if price_type == "retail" else "Giá đại lý",
        retail_visible=visible,
        incoming_products=incoming_products,
    )


@app.route("/")
def catalog():
    return render_catalog_page("dealer")


@app.route("/gia-le")
def retail_catalog():
    return render_catalog_page("retail")


@app.route("/api/retail-status")
def api_retail_status():
    db = SessionLocal()
    try:
        visible = retail_prices_visible(db)
    finally:
        db.close()
    return jsonify({"visible": visible})


@app.route("/api/products")
@app.route("/api/products/<price_type>")
def api_products(price_type="dealer"):
    """
    API cho trang khách.

    Điểm quan trọng:
    - Admin vẫn quản lý từng mã hàng riêng.
    - API trả model đã chuẩn hóa để frontend tự gộp chung card.
    - original_model giữ lại tên gốc nếu sau này cần xem/debug.
    """
    if price_type not in ("dealer", "retail"):
        return jsonify({"error": "Loại giá không hợp lệ"}), 404

    db = SessionLocal()
    try:
        visible = retail_prices_visible(db) if price_type == "retail" else True
        if price_type == "retail" and not visible:
            response = jsonify([])
            response.headers["X-Retail-Prices-Visible"] = "0"
            return response
        products = db.query(Product).order_by(Product.model, Product.color).all()

        out = []

        for p in products:
            d = p.to_dict()
            if price_type == "retail":
                d["price"] = p.retail_price if visible else None

            d["original_model"] = p.model
            d["model"] = normalize_model_name(p.model)

            d["image_url"] = p.image_filename if p.image_filename else None

            out.append(d)

        response = jsonify(out)
        response.headers["X-Retail-Prices-Visible"] = "1" if visible else "0"
        return response

    finally:
        db.close()


# ---------- Admin auth ----------
@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "").strip()

        if username == ADMIN_USERNAME and password == ADMIN_PASSWORD:
            login_user(AdminUser())
            return redirect(url_for("admin_dashboard"))

        flash("Sai tài khoản hoặc mật khẩu.")

    return render_template("login.html")


@app.route("/admin/logout")
@login_required
def admin_logout():
    logout_user()
    return redirect(url_for("catalog"))


# ---------- Admin dashboard ----------
@app.route("/admin/purchases/new", methods=["GET", "POST"])
@login_required
def admin_purchase_new():
    db = SessionLocal()
    try:
        if request.method == "POST":
            supplier_name = request.form.get("supplier_name", "").strip()
            supplier_code = request.form.get("supplier_code", "").strip()
            supplier_phone = request.form.get("supplier_phone", "").strip()
            supplier_email = request.form.get("supplier_email", "").strip()
            supplier_address = request.form.get("supplier_address", "").strip()
            receipt_number = request.form.get("receipt_number", "").strip()
            received_date = request.form.get("received_date", "").strip()
            expected_arrival_date_value = request.form.get("expected_arrival_date", "").strip()
            notes = request.form.get("notes", "").strip()

            try:
                items = json.loads(request.form.get("items_json", "[]"))
            except (TypeError, json.JSONDecodeError):
                items = []

            if not supplier_name:
                flash("Vui lòng nhập tên nhà cung cấp.")
                return redirect(url_for("admin_purchase_new"))
            if not receipt_number:
                flash("Vui lòng nhập mã phiếu nhập.")
                return redirect(url_for("admin_purchase_new"))
            if db.query(PurchaseReceipt).filter_by(receipt_number=receipt_number).first():
                flash(f"Mã phiếu '{receipt_number}' đã tồn tại.")
                return redirect(url_for("admin_purchase_new"))
            if not items:
                flash("Phiếu nhập cần có ít nhất một sản phẩm.")
                return redirect(url_for("admin_purchase_new"))

            supplier = None
            if supplier_code:
                supplier = db.query(Supplier).filter_by(code=supplier_code).first()
            if supplier is None:
                generated_code = supplier_code or f"NCC{(db.query(Supplier).count() + 1):04d}"
                supplier = Supplier(code=generated_code, name=supplier_name)
                db.add(supplier)
                db.flush()
            supplier.name = supplier_name
            supplier.phone = supplier_phone
            supplier.email = supplier_email
            supplier.address = supplier_address

            try:
                received_at = datetime.fromisoformat(received_date) if received_date else datetime.utcnow()
            except ValueError:
                received_at = datetime.utcnow()

            try:
                expected_arrival_date = datetime.fromisoformat(expected_arrival_date_value).date() if expected_arrival_date_value else None
            except ValueError:
                expected_arrival_date = None

            receipt = PurchaseReceipt(
                receipt_number=receipt_number,
                supplier_id=supplier.id,
                supplier_name=supplier.name,
                received_at=received_at,
                expected_arrival_date=expected_arrival_date,
                notes=notes,
                total_amount=0,
            )
            db.add(receipt)
            db.flush()

            total_amount = 0.0
            valid_lines = 0
            for raw_item in items:
                product_id = safe_int(raw_item.get("product_id"))
                quantity = max(0, safe_int(raw_item.get("quantity")))
                unit_price = max(0.0, safe_float(raw_item.get("unit_price")) or 0.0)
                product = db.query(Product).get(product_id) if product_id > 0 else None
                product_code = (product.code if product else str(raw_item.get("code", "")).strip())
                product_name = (product.model if product else str(raw_item.get("name", "")).strip())
                color = (product.color if product else str(raw_item.get("color", "")).strip())
                unit = (product.unit if product else str(raw_item.get("unit", "Chiếc")).strip()) or "Chiếc"
                if quantity <= 0 or not product_code or not product_name:
                    continue
                line_total = quantity * unit_price
                db.add(PurchaseReceiptItem(
                    receipt_id=receipt.id,
                    product_id=product.id if product else 0,
                    product_code=product_code,
                    product_name=product_name,
                    color=color or "",
                    unit=unit,
                    quantity=quantity,
                    unit_price=unit_price,
                    line_total=line_total,
                ))
                total_amount += line_total
                valid_lines += 1

            if valid_lines == 0:
                db.rollback()
                flash("Không có dòng sản phẩm hợp lệ trong phiếu nhập.")
                return redirect(url_for("admin_purchase_new"))

            receipt.total_amount = total_amount
            db.commit()
            flash(f"Đã lưu phiếu {receipt_number} với {valid_lines} sản phẩm. Tồn kho hiện tại không bị thay đổi.")
            return redirect(url_for("admin_purchase_new"))

        products = db.query(Product).order_by(Product.model, Product.color, Product.code).all()
        product_options = [
            {
                "id": p.id,
                "code": p.code,
                "name": p.model,
                "color": p.color or "",
                "unit": p.unit or "Chiếc",
                "price": p.price or 0,
            }
            for p in products
        ]
        suppliers = db.query(Supplier).order_by(Supplier.name).all()
        receipt_q = request.args.get("receipt_q", "").strip()
        supplier_q = request.args.get("supplier_q", "").strip()
        date_from = request.args.get("date_from", "").strip()
        date_to = request.args.get("date_to", "").strip()
        try:
            receipt_page = max(1, int(request.args.get("receipt_page", "1")))
        except ValueError:
            receipt_page = 1
        receipts_per_page = 20
        receipt_query = db.query(PurchaseReceipt)
        if receipt_q:
            receipt_query = receipt_query.filter(PurchaseReceipt.receipt_number.ilike(f"%{receipt_q}%"))
        if supplier_q:
            receipt_query = receipt_query.filter(PurchaseReceipt.supplier_name.ilike(f"%{supplier_q}%"))
        try:
            if date_from:
                receipt_query = receipt_query.filter(PurchaseReceipt.received_at >= datetime.fromisoformat(date_from))
            if date_to:
                receipt_query = receipt_query.filter(PurchaseReceipt.received_at < datetime.fromisoformat(date_to) + timedelta(days=1))
        except ValueError:
            date_from = ""
            date_to = ""
        receipt_total = receipt_query.count()
        receipt_total_pages = max(1, (receipt_total + receipts_per_page - 1) // receipts_per_page)
        if receipt_page > receipt_total_pages:
            receipt_page = receipt_total_pages
        recent_receipts = (
            receipt_query.order_by(PurchaseReceipt.received_at.desc(), PurchaseReceipt.id.desc())
            .offset((receipt_page - 1) * receipts_per_page)
            .limit(receipts_per_page)
            .all()
        )
        next_receipt_number = datetime.now().strftime("PN%Y%m%d-%H%M%S")
        return render_template(
            "purchase_form.html",
            products=products,
            product_options=product_options,
            suppliers=suppliers,
            recent_receipts=recent_receipts,
            receipt_q=receipt_q,
            supplier_q=supplier_q,
            date_from=date_from,
            date_to=date_to,
            receipt_page=receipt_page,
            receipt_total=receipt_total,
            receipt_total_pages=receipt_total_pages,
            next_receipt_number=next_receipt_number,
            today=datetime.now().strftime("%Y-%m-%d"),
        )
    finally:
        db.close()


@app.route("/admin/purchases/<int:receipt_id>")
@login_required
def admin_purchase_detail(receipt_id):
    db = SessionLocal()
    try:
        receipt = db.query(PurchaseReceipt).get(receipt_id)
        if not receipt:
            flash("Không tìm thấy phiếu nhập.")
            return redirect(url_for("admin_purchase_new"))
        supplier = db.query(Supplier).get(receipt.supplier_id) if receipt.supplier_id else None
        items = db.query(PurchaseReceiptItem).filter_by(receipt_id=receipt.id).order_by(PurchaseReceiptItem.id).all()
        return render_template("purchase_detail.html", receipt=receipt, supplier=supplier, items=items)
    finally:
        db.close()


@app.route("/admin/purchases/<int:receipt_id>/edit", methods=["GET", "POST"])
@login_required
def admin_purchase_edit(receipt_id):
    db = SessionLocal()
    try:
        receipt = db.query(PurchaseReceipt).get(receipt_id)
        if not receipt:
            flash("Không tìm thấy phiếu nhập.")
            return redirect(url_for("admin_purchase_new"))

        if request.method == "POST":
            supplier_name = request.form.get("supplier_name", "").strip()
            receipt_number = request.form.get("receipt_number", "").strip()
            try:
                submitted_items = json.loads(request.form.get("items_json", "[]"))
            except (TypeError, json.JSONDecodeError):
                submitted_items = []
            duplicate = db.query(PurchaseReceipt).filter(
                PurchaseReceipt.receipt_number == receipt_number,
                PurchaseReceipt.id != receipt.id,
            ).first()
            if not supplier_name or not receipt_number or not submitted_items or duplicate:
                flash("Vui lòng kiểm tra nhà cung cấp, mã phiếu và danh sách sản phẩm. Mã phiếu không được trùng.")
                return redirect(url_for("admin_purchase_edit", receipt_id=receipt.id))

            supplier_code = request.form.get("supplier_code", "").strip()
            supplier = db.query(Supplier).filter_by(code=supplier_code).first() if supplier_code else None
            if supplier is None:
                supplier = Supplier(code=supplier_code or f"NCC{(db.query(Supplier).count() + 1):04d}", name=supplier_name)
                db.add(supplier)
                db.flush()
            supplier.name = supplier_name
            supplier.phone = request.form.get("supplier_phone", "").strip()
            supplier.email = request.form.get("supplier_email", "").strip()
            supplier.address = request.form.get("supplier_address", "").strip()

            try:
                received_at = datetime.fromisoformat(request.form.get("received_date", ""))
            except ValueError:
                received_at = receipt.received_at
            expected_arrival_date_value = request.form.get("expected_arrival_date", "").strip()
            try:
                expected_arrival_date = datetime.fromisoformat(expected_arrival_date_value).date() if expected_arrival_date_value else None
            except ValueError:
                expected_arrival_date = receipt.expected_arrival_date
            receipt.receipt_number = receipt_number
            receipt.supplier_id = supplier.id
            receipt.supplier_name = supplier.name
            receipt.received_at = received_at
            receipt.expected_arrival_date = expected_arrival_date
            receipt.notes = request.form.get("notes", "").strip()
            db.query(PurchaseReceiptItem).filter_by(receipt_id=receipt.id).delete(synchronize_session=False)

            total_amount = 0.0
            valid_lines = 0
            for raw_item in submitted_items:
                product_id = safe_int(raw_item.get("product_id"))
                product = db.query(Product).get(product_id) if product_id > 0 else None
                quantity = max(0, safe_int(raw_item.get("quantity")))
                unit_price = max(0.0, safe_float(raw_item.get("unit_price")) or 0.0)
                product_code = product.code if product else str(raw_item.get("code", "")).strip()
                product_name = product.model if product else str(raw_item.get("name", "")).strip()
                color = product.color if product else str(raw_item.get("color", "")).strip()
                unit = (product.unit if product else str(raw_item.get("unit", "Chiếc")).strip()) or "Chiếc"
                if quantity <= 0 or not product_code or not product_name:
                    continue
                line_total = quantity * unit_price
                db.add(PurchaseReceiptItem(receipt_id=receipt.id, product_id=product.id if product else 0, product_code=product_code, product_name=product_name, color=color or "", unit=unit, quantity=quantity, unit_price=unit_price, line_total=line_total))
                total_amount += line_total
                valid_lines += 1
            if valid_lines == 0:
                db.rollback()
                flash("Phiếu nhập cần có ít nhất một dòng sản phẩm hợp lệ.")
                return redirect(url_for("admin_purchase_edit", receipt_id=receipt.id))
            receipt.total_amount = total_amount
            db.commit()
            flash(f"Đã cập nhật phiếu {receipt.receipt_number}. Tồn kho hiện tại không bị thay đổi.")
            return redirect(url_for("admin_purchase_detail", receipt_id=receipt.id))

        products = db.query(Product).order_by(Product.model, Product.color, Product.code).all()
        product_options = [{"id": p.id, "code": p.code, "name": p.model, "color": p.color or "", "unit": p.unit or "Chiếc", "price": p.price or 0} for p in products]
        suppliers = db.query(Supplier).order_by(Supplier.name).all()
        edit_supplier = db.query(Supplier).get(receipt.supplier_id) if receipt.supplier_id else None
        stored_items = db.query(PurchaseReceiptItem).filter_by(receipt_id=receipt.id).order_by(PurchaseReceiptItem.id).all()
        initial_items = [{"product_id": item.product_id, "code": item.product_code, "name": item.product_name, "color": item.color or "", "unit": item.unit or "Chiếc", "quantity": item.quantity, "unit_price": item.unit_price or 0} for item in stored_items]
        return render_template("purchase_form.html", products=products, product_options=product_options, suppliers=suppliers, edit_receipt=receipt, edit_supplier=edit_supplier, initial_items=initial_items, next_receipt_number=receipt.receipt_number, today=receipt.received_at.strftime("%Y-%m-%d"))
    finally:
        db.close()


@app.route("/admin/purchases/<int:receipt_id>/delete", methods=["POST"])
@login_required
def admin_purchase_delete(receipt_id):
    db = SessionLocal()
    try:
        receipt = db.query(PurchaseReceipt).get(receipt_id)
        if not receipt:
            flash("Không tìm thấy phiếu nhập.")
        else:
            receipt_number = receipt.receipt_number
            db.query(PurchaseReceiptItem).filter_by(receipt_id=receipt.id).delete(synchronize_session=False)
            db.delete(receipt)
            db.commit()
            flash(f"Đã xóa phiếu {receipt_number}. Tồn kho hiện tại không bị thay đổi.")
    finally:
        db.close()
    return redirect(url_for("admin_purchase_history"))


def purchase_history_query(db):
    query = db.query(PurchaseReceipt)
    keyword = request.args.get("q", "").strip()
    supplier_id = request.args.get("supplier_id", "").strip()
    month = request.args.get("month", "").strip()
    year = request.args.get("year", "").strip()
    date_from = request.args.get("date_from", "").strip()
    date_to = request.args.get("date_to", "").strip()
    if keyword:
        like = f"%{keyword}%"
        item_receipts = db.query(PurchaseReceiptItem.receipt_id).filter(
            or_(PurchaseReceiptItem.product_name.ilike(like), PurchaseReceiptItem.product_code.ilike(like))
        )
        query = query.filter(or_(PurchaseReceipt.receipt_number.ilike(like), PurchaseReceipt.supplier_name.ilike(like), PurchaseReceipt.id.in_(item_receipts)))
    if supplier_id.isdigit():
        query = query.filter(PurchaseReceipt.supplier_id == int(supplier_id))
    try:
        if month:
            query = query.filter(func.extract("month", PurchaseReceipt.received_at) == int(month))
        if year:
            query = query.filter(func.extract("year", PurchaseReceipt.received_at) == int(year))
        if date_from:
            query = query.filter(PurchaseReceipt.received_at >= datetime.fromisoformat(date_from))
        if date_to:
            query = query.filter(PurchaseReceipt.received_at < datetime.fromisoformat(date_to) + timedelta(days=1))
    except ValueError:
        pass
    return query


@app.route("/admin/purchases")
@login_required
def admin_purchase_history():
    db = SessionLocal()
    try:
        query = purchase_history_query(db)
        try:
            page = max(1, int(request.args.get("page", "1")))
        except ValueError:
            page = 1
        try:
            per_page = int(request.args.get("per_page", "10"))
        except ValueError:
            per_page = 10
        if per_page not in (10, 20, 50):
            per_page = 10
        total_receipts = query.count()
        total_pages = max(1, (total_receipts + per_page - 1) // per_page)
        page = min(page, total_pages)
        receipts = query.order_by(PurchaseReceipt.received_at.desc(), PurchaseReceipt.id.desc()).offset((page - 1) * per_page).limit(per_page).all()
        all_receipts = query.all()
        receipt_ids = [receipt.id for receipt in all_receipts]
        total_amount = sum((receipt.total_amount or 0) for receipt in all_receipts)
        supplier_count = len({receipt.supplier_id or receipt.supplier_name for receipt in all_receipts})
        item_totals = {}
        total_quantity = 0
        if receipt_ids:
            grouped_items = db.query(PurchaseReceiptItem.receipt_id, func.count(PurchaseReceiptItem.id), func.sum(PurchaseReceiptItem.quantity)).filter(PurchaseReceiptItem.receipt_id.in_(receipt_ids)).group_by(PurchaseReceiptItem.receipt_id).all()
            item_totals = {rid: {"lines": int(lines or 0), "quantity": int(quantity or 0)} for rid, lines, quantity in grouped_items}
            total_quantity = sum(value["quantity"] for value in item_totals.values())
        supplier_summary = {}
        for receipt in all_receipts:
            key = receipt.supplier_name
            row = supplier_summary.setdefault(key, {"name": key, "receipts": 0, "lines": 0, "quantity": 0, "amount": 0})
            row["receipts"] += 1
            row["lines"] += item_totals.get(receipt.id, {}).get("lines", 0)
            row["quantity"] += item_totals.get(receipt.id, {}).get("quantity", 0)
            row["amount"] += receipt.total_amount or 0
        supplier_summary = sorted(supplier_summary.values(), key=lambda row: row["amount"], reverse=True)
        chart_summary = supplier_summary[:5]
        if len(supplier_summary) > 5:
            chart_summary = chart_summary + [{"name": "Khác", "amount": sum(row["amount"] for row in supplier_summary[5:])}]
        chart_labels = [row["name"] for row in chart_summary]
        chart_values = [row["amount"] for row in chart_summary]
        suppliers = db.query(Supplier).order_by(Supplier.name).all()
        years = sorted({receipt.received_at.year for receipt in db.query(PurchaseReceipt).all() if receipt.received_at}, reverse=True)
        filter_args = {key: value for key, value in request.args.items() if key not in ("page", "per_page") and value}
        return render_template("purchase_history.html", receipts=receipts, item_totals=item_totals, supplier_summary=supplier_summary, chart_labels=chart_labels, chart_values=chart_values, suppliers=suppliers, years=years, total_receipts=total_receipts, total_amount=total_amount, total_quantity=total_quantity, supplier_count=supplier_count, page=page, per_page=per_page, total_pages=total_pages, filters=request.args, filter_args=filter_args)
    finally:
        db.close()


@app.route("/admin/purchases/export")
@login_required
def admin_purchase_export():
    db = SessionLocal()
    try:
        receipts = purchase_history_query(db).order_by(PurchaseReceipt.received_at.desc()).all()
        rows = []
        for receipt in receipts:
            items = db.query(PurchaseReceiptItem).filter_by(receipt_id=receipt.id).all()
            for item in items:
                rows.append({"Mã phiếu": receipt.receipt_number, "Ngày nhập": receipt.received_at.strftime("%d/%m/%Y"), "Ngày dự kiến": receipt.expected_arrival_date.strftime("%d/%m/%Y") if receipt.expected_arrival_date else "", "Nhà cung cấp": receipt.supplier_name, "Mã hàng": item.product_code, "Sản phẩm": item.product_name, "Màu": item.color, "Đơn vị": item.unit, "Số lượng": item.quantity, "Giá nhập": item.unit_price, "Thành tiền": item.line_total, "Ghi chú": receipt.notes})
        output = BytesIO()
        pd.DataFrame(rows).to_excel(output, index=False, sheet_name="Lich su nhap hang")
        output.seek(0)
        return send_file(output, as_attachment=True, download_name=f"lich_su_nhap_hang_{datetime.now():%Y%m%d}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    finally:
        db.close()


@app.route("/admin/suppliers")
@login_required
def admin_suppliers():
    db = SessionLocal()
    try:
        q = request.args.get("q", "").strip()
        supplier_query = db.query(Supplier).join(PurchaseReceipt, PurchaseReceipt.supplier_id == Supplier.id).distinct()
        if q:
            like = f"%{q}%"
            supplier_query = supplier_query.filter(or_(Supplier.name.ilike(like), Supplier.code.ilike(like), Supplier.phone.ilike(like), Supplier.email.ilike(like)))
        suppliers = supplier_query.order_by(Supplier.name).all()
        supplier_rows = []
        all_product_codes = set()
        total_value = 0.0
        total_quantity = 0
        for supplier in suppliers:
            receipts = db.query(PurchaseReceipt).filter_by(supplier_id=supplier.id).order_by(PurchaseReceipt.received_at.desc()).all()
            receipt_ids = [receipt.id for receipt in receipts]
            items = db.query(PurchaseReceiptItem).filter(PurchaseReceiptItem.receipt_id.in_(receipt_ids)).all() if receipt_ids else []
            product_codes = {item.product_code for item in items}
            value = sum((receipt.total_amount or 0) for receipt in receipts)
            quantity = sum((item.quantity or 0) for item in items)
            all_product_codes.update(product_codes)
            total_value += value
            total_quantity += quantity
            supplier_rows.append({"supplier": supplier, "receipt_count": len(receipts), "product_count": len(product_codes), "quantity": quantity, "value": value, "last_received": receipts[0].received_at if receipts else None})
        supplier_rows.sort(key=lambda row: (row["last_received"] is not None, row["last_received"] or datetime.min), reverse=True)
        return render_template("suppliers.html", supplier_rows=supplier_rows, q=q, total_suppliers=len(supplier_rows), total_products=len(all_product_codes), total_value=total_value, total_quantity=total_quantity)
    finally:
        db.close()


@app.route("/admin/suppliers/<int:supplier_id>")
@login_required
def admin_supplier_detail(supplier_id):
    db = SessionLocal()
    try:
        supplier = db.query(Supplier).get(supplier_id)
        if not supplier:
            flash("Không tìm thấy nhà cung cấp.")
            return redirect(url_for("admin_suppliers"))
        receipts = db.query(PurchaseReceipt).filter_by(supplier_id=supplier.id).order_by(PurchaseReceipt.received_at.desc()).all()
        receipt_ids = [receipt.id for receipt in receipts]
        records = []
        if receipt_ids:
            records = db.query(PurchaseReceiptItem, PurchaseReceipt).join(PurchaseReceipt, PurchaseReceipt.id == PurchaseReceiptItem.receipt_id).filter(PurchaseReceiptItem.receipt_id.in_(receipt_ids)).order_by(PurchaseReceipt.received_at.desc(), PurchaseReceiptItem.id.desc()).all()
        grouped = {}
        for item, receipt in records:
            key = (item.product_code, item.product_name, item.color or "")
            row = grouped.setdefault(key, {"code": item.product_code, "name": item.product_name, "color": item.color or "", "prices": [], "latest_price": item.unit_price or 0, "latest_date": receipt.received_at, "times": 0, "quantity": 0})
            row["prices"].append(item.unit_price or 0)
            row["times"] += 1
            row["quantity"] += item.quantity or 0
        price_rows = []
        for row in grouped.values():
            positive_prices = [price for price in row["prices"] if price > 0]
            row["min_price"] = min(positive_prices) if positive_prices else 0
            row["max_price"] = max(positive_prices) if positive_prices else 0
            row["avg_price"] = sum(positive_prices) / len(positive_prices) if positive_prices else 0
            price_rows.append(row)
        price_rows.sort(key=lambda row: (row["name"], row["color"], row["code"]))
        total_value = sum((receipt.total_amount or 0) for receipt in receipts)
        total_quantity = sum((item.quantity or 0) for item, _ in records)
        return render_template("supplier_detail.html", supplier=supplier, receipts=receipts[:20], price_rows=price_rows, receipt_count=len(receipts), product_count=len(price_rows), total_value=total_value, total_quantity=total_quantity, first_received=receipts[-1].received_at if receipts else None, last_received=receipts[0].received_at if receipts else None)
    finally:
        db.close()


@app.route("/admin/suppliers/catalog")
@login_required
def admin_supplier_catalog():
    db = SessionLocal()
    try:
        q = request.args.get("q", "").strip().lower()
        suppliers = db.query(Supplier).join(PurchaseReceipt, PurchaseReceipt.supplier_id == Supplier.id).distinct().order_by(Supplier.name).all()
        supplier_catalog = []
        supplier_colors = {supplier.id: index % 10 for index, supplier in enumerate(sorted(suppliers, key=lambda item: item.id))}
        all_codes = set()
        total_links = 0
        for supplier in suppliers:
            records = (
                db.query(PurchaseReceiptItem, PurchaseReceipt)
                .join(PurchaseReceipt, PurchaseReceipt.id == PurchaseReceiptItem.receipt_id)
                .filter(PurchaseReceipt.supplier_id == supplier.id)
                .order_by(PurchaseReceipt.received_at.desc(), PurchaseReceiptItem.id.desc())
                .all()
            )
            unique_products = {}
            for item, receipt in records:
                if item.product_code in unique_products:
                    continue
                unique_products[item.product_code] = {
                    "code": item.product_code,
                    "name": item.product_name,
                    "color": item.color or "",
                    "latest_price": item.unit_price or 0,
                    "latest_date": receipt.received_at,
                }
            products = list(unique_products.values())
            if q:
                supplier_matches = q in supplier.name.lower() or q in supplier.code.lower()
                matched_products = [product for product in products if q in f"{product['code']} {product['name']} {product['color']}".lower()]
                if not supplier_matches and not matched_products:
                    continue
                if not supplier_matches:
                    products = matched_products
            if not products and q:
                continue
            products.sort(key=lambda product: (product["name"], product["color"], product["code"]))
            all_codes.update(product["code"] for product in products)
            total_links += len(products)
            supplier_catalog.append({"supplier": supplier, "products": products, "product_count": len(products), "color_index": supplier_colors[supplier.id]})
        supplier_catalog.sort(key=lambda row: (-row["product_count"], row["supplier"].name))
        return render_template("supplier_catalog.html", supplier_catalog=supplier_catalog, q=request.args.get("q", "").strip(), supplier_count=len(supplier_catalog), unique_product_count=len(all_codes), total_links=total_links)
    finally:
        db.close()


@app.route("/admin/suppliers/compare")
@login_required
def admin_supplier_compare():
    db = SessionLocal()
    try:
        q = request.args.get("q", "").strip()
        search_key = product_match_key(q) if q else ""
        records = db.query(PurchaseReceiptItem, PurchaseReceipt).join(PurchaseReceipt, PurchaseReceipt.id == PurchaseReceiptItem.receipt_id).order_by(PurchaseReceipt.received_at.desc()).all()
        grouped = {}
        for item, receipt in records:
            if search_key:
                searchable = (item.product_code, item.product_name, item.color)
                if not any(search_key in product_match_key(value) for value in searchable):
                    continue
            key = (receipt.supplier_id, receipt.supplier_name, item.product_code, item.product_name, item.color or "")
            row = grouped.setdefault(key, {"supplier_id": receipt.supplier_id, "supplier_name": receipt.supplier_name, "code": item.product_code, "name": item.product_name, "color": item.color or "", "prices": [], "latest_price": item.unit_price or 0, "latest_date": receipt.received_at, "times": 0, "quantity": 0})
            row["prices"].append(item.unit_price or 0)
            row["times"] += 1
            row["quantity"] += item.quantity or 0
        compare_rows = list(grouped.values())
        for row in compare_rows:
            positive = [price for price in row["prices"] if price > 0]
            row["min_price"] = min(positive) if positive else 0
            row["max_price"] = max(positive) if positive else 0
            row["avg_price"] = sum(positive) / len(positive) if positive else 0
        compare_rows.sort(key=lambda row: (row["code"], row["latest_price"] or float("inf"), row["supplier_name"]))
        cheapest_by_code = {}
        for row in compare_rows:
            if row["latest_price"] > 0:
                cheapest_by_code[row["code"]] = min(cheapest_by_code.get(row["code"], row["latest_price"]), row["latest_price"])
        for row in compare_rows:
            row["is_cheapest"] = row["latest_price"] > 0 and row["latest_price"] == cheapest_by_code.get(row["code"])
            row["age_days"] = (datetime.now() - row["latest_date"]).days if row["latest_date"] else None
        return render_template("supplier_compare.html", q=q, compare_rows=compare_rows)
    finally:
        db.close()


@app.route("/admin")
@login_required
def admin_dashboard():
    page = request.args.get("page", "1")
    q = request.args.get("q", "").strip()
    price_type = request.args.get("price_type", "dealer")
    if price_type not in ("dealer", "retail"):
        price_type = "dealer"

    try:
        page = max(1, int(page))
    except ValueError:
        page = 1

    per_page = 10

    db = SessionLocal()
    try:
        base_query = db.query(Product)

        if q:
            like = f"%{q}%"
            base_query = base_query.filter(
                (Product.code.ilike(like)) |
                (Product.model.ilike(like)) |
                (Product.color.ilike(like)) |
                (Product.category.ilike(like))
            )

        filtered_total = base_query.count()

        products = (
            base_query
            .order_by(Product.model, Product.color, Product.code)
            .offset((page - 1) * per_page)
            .limit(per_page)
            .all()
        )

        all_products = db.query(Product).all()

        categories = sorted({
            p.category for p in all_products
            if p.category and str(p.category).strip()
        })

        total_products = len(all_products)
        total_stock = db.query(Product).with_entities(func.sum(Product.stock)).scalar() or 0
        total_available = sum([p.available for p in all_products])
        low_stock_products = sorted(
            [p for p in all_products if p.available <= 5],
            key=lambda p: (p.available, p.model or "", p.code or ""),
        )[:5]
        low_stock_count = sum(1 for p in all_products if p.available <= 5)
        recent_products = sorted(
            [p for p in all_products if p.updated_at],
            key=lambda p: p.updated_at,
            reverse=True,
        )[:5]

        total_pages = max(1, (filtered_total + per_page - 1) // per_page)
        retail_visible = retail_prices_visible(db)
        retail_priced_count = db.query(Product).filter(Product.retail_price > 0).count()

        return render_template(
            "dashboard.html",
            products=products,
            categories=categories,
            page=page,
            total_pages=total_pages,
            total_products=total_products,
            total_stock=total_stock,
            total_available=total_available,
            low_stock_products=low_stock_products,
            low_stock_count=low_stock_count,
            recent_products=recent_products,
            filtered_total=filtered_total,
            per_page=per_page,
            q=q,
            price_type=price_type,
            retail_visible=retail_visible,
            retail_priced_count=retail_priced_count,
        )

    finally:
        db.close()


@app.route("/api/reorder-suggestions")
def api_reorder_suggestions():
    """Suggest replenishment when stock is below two months of average sales."""
    try:
        monthly_sales = load_monthly_sales()
    except Exception as error:
        print(f"Khong doc duoc file lich su ban hang: {error}")
        return jsonify({"items": [], "matched_models": 0, "error": "sales_file_unavailable"}), 500

    db = SessionLocal()
    try:
        latest_inventory_update = db.query(func.max(Product.updated_at)).scalar()
        grouped = {}
        for product in db.query(Product).order_by(Product.model, Product.color).all():
            key = product_match_key(product.model)
            if not key:
                continue

            if key not in grouped:
                grouped[key] = {
                    "model": normalize_model_name(product.model),
                    "category": product.category or "",
                    "stock": 0,
                    "image_url": None,
                }

            grouped[key]["stock"] += max(product.stock or 0, 0)
            if not grouped[key]["image_url"] and product.image_filename:
                grouped[key]["image_url"] = product.image_filename

        suggestions = []
        matched_models = 0
        for key, product in grouped.items():
            average_monthly = monthly_sales.get(key)
            if average_monthly is None:
                continue

            matched_models += 1
            reserve_exact = average_monthly * 2
            if product["stock"] >= reserve_exact:
                continue

            suggestions.append({
                **product,
                "average_monthly_sales": round(average_monthly, 1),
                "reserve_target": math.ceil(reserve_exact),
                "reorder_quantity": math.ceil(reserve_exact - product["stock"]),
            })

        suggestions.sort(
            key=lambda item: (item["reorder_quantity"], item["average_monthly_sales"]),
            reverse=True,
        )
        return jsonify({
            "items": suggestions,
            "matched_models": matched_models,
            "reserve_months": 2,
            "inventory_updated_at": (
                latest_inventory_update.isoformat()
                + ("Z" if latest_inventory_update.tzinfo is None else "")
                if latest_inventory_update else None
            ),
        })
    finally:
        db.close()


@app.route("/admin/retail-visibility", methods=["POST"])
@login_required
def admin_retail_visibility():
    enabled = request.form.get("enabled")
    if enabled not in ("0", "1"):
        flash("Trạng thái hiển thị không hợp lệ.")
        return redirect(url_for("admin_dashboard", price_type="retail"))

    db = SessionLocal()
    try:
        setting = db.query(AppSetting).filter_by(key="retail_prices_visible").first()
        if setting is None:
            setting = AppSetting(key="retail_prices_visible", value=enabled)
            db.add(setting)
        else:
            setting.value = enabled
        db.commit()
        flash("Đã bật module Giá lẻ." if enabled == "1" else "Đã tạm ẩn module Giá lẻ.")
    finally:
        db.close()
    return redirect(url_for("admin_dashboard", price_type="retail"))


@app.route("/admin/products/add", methods=["POST"])
@login_required
def admin_add_product():
    db = SessionLocal()
    try:
        code = request.form.get("code", "").strip()

        if not code:
            flash("Vui lòng nhập mã hàng.")
            return redirect(url_for("admin_dashboard"))

        if db.query(Product).filter_by(code=code).first():
            flash(f"Mã hàng '{code}' đã tồn tại.")
            return redirect(url_for("admin_dashboard"))

        model = request.form.get("model", "").strip()

        if not model:
            flash("Vui lòng nhập tên mẫu xe.")
            return redirect(url_for("admin_dashboard"))

        price = request.form.get("price", "").strip()

        p = Product(
            code=code,
            model=normalize_model_name(model),
            color=request.form.get("color", "").strip(),
            category=request.form.get("category", "").strip(),
            unit=request.form.get("unit", "Chiếc").strip() or "Chiếc",
            stock=safe_int(request.form.get("stock")),
            reserved=safe_int(request.form.get("reserved")),
            price=safe_float(price),
            retail_price=safe_float(request.form.get("retail_price", "").strip()),
        )

        db.add(p)
        db.commit()

        file = request.files.get("image")
        image_name = save_product_image(file, p.id)

        if image_name:
            p.image_filename = image_name
            db.commit()

        flash("Đã thêm sản phẩm.")

    finally:
        db.close()

    return redirect(url_for("admin_dashboard"))


@app.route("/admin/products/<int:product_id>/edit", methods=["POST"])
@login_required
def admin_edit_product(product_id):
    db = SessionLocal()
    try:
        p = db.query(Product).get(product_id)

        if not p:
            flash("Không tìm thấy sản phẩm.")
            return redirect(url_for("admin_dashboard"))

        new_code = request.form.get("code", p.code).strip()

        if not new_code:
            flash("Mã hàng không được để trống.")
            return redirect(url_for("admin_dashboard"))

        duplicate = (
            db.query(Product)
            .filter(Product.code == new_code, Product.id != product_id)
            .first()
        )

        if duplicate:
            flash(f"Mã hàng '{new_code}' đã tồn tại ở sản phẩm khác.")
            return redirect(url_for("admin_dashboard"))

        model = request.form.get("model", "").strip()

        if not model:
            flash("Tên mẫu xe không được để trống.")
            return redirect(url_for("admin_dashboard"))

        p.code = new_code
        p.model = normalize_model_name(model)
        p.color = request.form.get("color", "").strip()
        p.category = request.form.get("category", "").strip()
        p.unit = request.form.get("unit", "Chiếc").strip() or "Chiếc"
        p.stock = safe_int(request.form.get("stock"))
        p.reserved = safe_int(request.form.get("reserved"))
        p.price = safe_float(request.form.get("price", "").strip())
        p.retail_price = safe_float(request.form.get("retail_price", "").strip())

        file = request.files.get("image")
        image_name = save_product_image(file, p.id)

        if image_name:
            p.image_filename = image_name

        db.commit()

        flash("Đã cập nhật sản phẩm.")

    finally:
        db.close()

    return redirect(url_for("admin_dashboard"))


@app.route("/admin/products/<int:product_id>/delete-image", methods=["POST"])
@login_required
def admin_delete_product_image(product_id):
    db = SessionLocal()
    try:
        p = db.query(Product).get(product_id)

        if not p:
            flash("Không tìm thấy sản phẩm.")
            return redirect(url_for("admin_dashboard"))

        if p.image_filename:
            try:
                public_id = f"abraham_inventory/products/product_{p.id}"
                cloudinary.uploader.destroy(public_id, resource_type="image")
            except Exception as e:
                print(f"Loi xoa anh Cloudinary: {e}")

            p.image_filename = ""
            db.commit()
            flash("Đã xóa ảnh sản phẩm.")
        else:
            flash("Sản phẩm này chưa có ảnh.")

    finally:
        db.close()

    return redirect(url_for("admin_dashboard"))


@app.route("/admin/products/<int:product_id>/delete", methods=["POST"])
@login_required
def admin_delete_product(product_id):
    db = SessionLocal()
    try:
        p = db.query(Product).get(product_id)

        if p:
            db.delete(p)
            db.commit()
            flash("Đã xóa sản phẩm.")

    finally:
        db.close()

    return redirect(url_for("admin_dashboard"))


@app.route("/admin/import", methods=["POST"])
@login_required
def admin_import():
    file = request.files.get("file")

    if not file or file.filename == "":
        flash("Chưa chọn file để nhập.")
        return redirect(url_for("admin_dashboard"))

    try:
        if file.filename.lower().endswith(".csv"):
            df = pd.read_csv(file)
        else:
            df = pd.read_excel(file)
    except Exception as e:
        flash(f"Không đọc được file: {e}")
        return redirect(url_for("admin_dashboard"))

    if "Mã hàng hóa" not in df.columns:
        flash("File thiếu cột bắt buộc: Mã hàng hóa.")
        return redirect(url_for("admin_dashboard"))

    has_name = "Tên hàng hóa" in df.columns
    has_category = "Loại hàng hóa" in df.columns
    unit_col = "Đơn vị tính" if "Đơn vị tính" in df.columns else (
        "Đơn vị tính chính" if "Đơn vị tính chính" in df.columns else None
    )
    has_stock = "Số lượng tồn" in df.columns
    has_reserved = "SL đã đặt chưa giao" in df.columns
    has_price = "Đơn giá bán" in df.columns

    df = df.fillna(0)

    db = SessionLocal()
    added, updated, skipped = 0, 0, 0

    BATCH_SIZE = 1000

    try:
        # Lay toan bo (id, code) hien co 1 lan duy nhat, khong load ca
        # object de tiet kiem bo nho va thoi gian query.
        existing_by_code = {
            code: pid for pid, code in db.query(Product.id, Product.code).all()
        }

        to_insert = []
        to_update = []

        for row in df.itertuples(index=False):
            row = dict(zip(df.columns, row))

            code = str(row["Mã hàng hóa"]).strip()

            if not code or code == "0":
                continue

            existing_id = existing_by_code.get(code)

            if existing_id:
                # San pham da co: chi cap nhat cot nao co trong file,
                # cot khong co trong file thi giu nguyen du lieu cu.
                data = {"id": existing_id}

                if has_name:
                    name = str(row["Tên hàng hóa"]).strip()
                    if name and name != "0":
                        model, color = split_model_color(name)
                        data["model"] = model
                        data["color"] = color

                if has_category:
                    category = str(row.get("Loại hàng hóa", "")).strip()
                    if category and category != "0":
                        data["category"] = category

                if unit_col:
                    unit = str(row.get(unit_col, "")).strip()
                    if unit and unit != "0":
                        data["unit"] = unit

                if has_stock:
                    data["stock"] = safe_int(row.get("Số lượng tồn", 0))

                if has_reserved:
                    data["reserved"] = safe_int(row.get("SL đã đặt chưa giao", 0))

                if has_price:
                    price_value = safe_float(row.get("Đơn giá bán"))
                    if price_value is not None:
                        data["price"] = price_value

                to_update.append(data)
                updated += 1

            else:
                # San pham chua co: bat buoc phai co ten de tao moi.
                if not has_name:
                    skipped += 1
                    continue

                name = str(row["Tên hàng hóa"]).strip()

                if not name or name == "0":
                    skipped += 1
                    continue

                model, color = split_model_color(name)

                to_insert.append({
                    "code": code,
                    "model": model,
                    "color": color,
                    "category": str(row.get("Loại hàng hóa", "")).strip() if has_category else "",
                    "unit": (str(row.get(unit_col, "Chiếc")).strip() or "Chiếc") if unit_col else "Chiếc",
                    "stock": safe_int(row.get("Số lượng tồn", 0)) if has_stock else 0,
                    "reserved": safe_int(row.get("SL đã đặt chưa giao", 0)) if has_reserved else 0,
                    "price": safe_float(row.get("Đơn giá bán")) if has_price else None,
                })
                added += 1

            # Flush theo batch de tranh giu transaction qua lon trong bo nho.
            if len(to_insert) >= BATCH_SIZE:
                db.bulk_insert_mappings(Product, to_insert)
                db.commit()
                to_insert.clear()

            if len(to_update) >= BATCH_SIZE:
                db.bulk_update_mappings(Product, to_update)
                db.commit()
                to_update.clear()

        # Flush phan con lai.
        if to_insert:
            db.bulk_insert_mappings(Product, to_insert)
            db.commit()

        if to_update:
            db.bulk_update_mappings(Product, to_update)
            db.commit()

        message = f"Nhập xong: thêm mới {added}, cập nhật {updated}."
        if skipped:
            message += f" Bỏ qua {skipped} dòng thiếu tên hàng."
        flash(message)

    except Exception as e:
        db.rollback()
        flash(f"Lỗi khi nhập dữ liệu: {e}")

    finally:
        db.close()

    return redirect(url_for("admin_dashboard"))


@app.route("/admin/import-price", methods=["POST"])
@login_required
def admin_import_price():
    """
    Import dealer or retail prices by product code. PostgreSQL batches updates;
    SQLite uses individual updates for compatibility.
    """
    price_type = request.form.get("price_type", "dealer")
    if price_type not in ("dealer", "retail"):
        flash("Loại giá không hợp lệ.")
        return redirect(url_for("admin_dashboard"))
    price_column = "retail_price" if price_type == "retail" else "price"
    price_label = "Giá lẻ" if price_type == "retail" else "Giá đại lý"

    file = request.files.get("file")

    if not file or file.filename == "":
        flash("Chưa chọn file để nhập giá.")
        return redirect(url_for("admin_dashboard", price_type=price_type))

    try:
        if file.filename.lower().endswith(".csv"):
            df = pd.read_csv(file)
        else:
            df = pd.read_excel(file)
    except Exception as e:
        flash(f"Không đọc được file: {e}")
        return redirect(url_for("admin_dashboard", price_type=price_type))

    if "Mã hàng hóa" not in df.columns or "Đơn giá bán" not in df.columns:
        flash("File thiếu cột bắt buộc: Mã hàng hóa hoặc Đơn giá bán.")
        return redirect(url_for("admin_dashboard", price_type=price_type))

    df = df[["Mã hàng hóa", "Đơn giá bán"]]

    rows = []
    skipped = 0

    for row in df.itertuples(index=False):
        code = str(row[0]).strip()
        price = safe_float(row[1])

        if not code or code == "0" or price is None:
            skipped += 1
            continue

        rows.append((code, price))

    if not rows:
        flash("Không có dòng nào hợp lệ để cập nhật giá.")
        return redirect(url_for("admin_dashboard", price_type=price_type))

    try:
        # Gop nhieu dong thanh 1 cau UPDATE duy nhat cho moi chunk,
        # thay vi executemany tung dong -> giam so round-trip toi DB.
        CHUNK = 2000
        matched = 0

        with engine.begin() as conn:
            for i in range(0, len(rows), CHUNK):
                chunk = rows[i:i + CHUNK]

                if engine.dialect.name == "sqlite":
                    statement = text(f"UPDATE products SET {price_column} = :price WHERE code = :code")
                    for code, price in chunk:
                        result = conn.execute(statement, {"code": code, "price": price})
                        matched += result.rowcount or 0
                    continue

                values_sql = ", ".join(
                    f"(:code{j}, :price{j})" for j in range(len(chunk))
                )
                params = {}
                for j, (code, price) in enumerate(chunk):
                    params[f"code{j}"] = code
                    params[f"price{j}"] = price

                sql = text(f"""
                    UPDATE products AS p
                    SET {price_column} = c.price
                    FROM (VALUES {values_sql}) AS c(code, price)
                    WHERE p.code = c.code
                """)

                result = conn.execute(sql, params)
                matched += result.rowcount or 0

        message = f"Đã cập nhật {price_label} cho {matched} mã hàng."
        if skipped:
            message += f" Bỏ qua {skipped} dòng thiếu mã hoặc giá."
        flash(message)

    except Exception as e:
        flash(f"Lỗi khi cập nhật giá: {e}")

    return redirect(url_for("admin_dashboard", price_type=price_type))


if __name__ == "__main__":
    seed_if_empty()
    app.run(
        host="0.0.0.0",
        port=5000,
        debug=os.environ.get("FLASK_DEBUG", "0") == "1",
    )
